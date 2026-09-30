from __future__ import annotations

from html import escape
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = ROOT / "docs" / "report" / "FITROOM_PROJECT_REPORT.docx"
PDF_PATH = ROOT / "output" / "pdf" / "FITROOM_PROJECT_REPORT.pdf"
ASSET_DIR = ROOT / "docs" / "report" / "assets"
FONT_PATH = Path("/System/Library/Fonts/Supplemental/AppleGothic.ttf")

BLUE = "2557D6"
DARK = "121826"
INK = "20242B"
MUTED = "667085"
PALE = "EEF3FF"
LINE = "DDE2EA"


def rl_color(value: str):
    return colors.HexColor(f"#{value}")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "AppleGothic"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "AppleGothic")
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "AppleGothic")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "AppleGothic")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def configure_docx_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "AppleGothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "AppleGothic")
    normal.font.size = Pt(9.3)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.2

    title = doc.styles["Title"]
    title.font.name = "AppleGothic"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "AppleGothic")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(DARK)
    title.paragraph_format.space_after = Pt(10)

    for name, size, color in (("Heading 1", 19, DARK), ("Heading 2", 13, BLUE), ("Heading 3", 10.5, DARK)):
        style = doc.styles[name]
        style.font.name = "AppleGothic"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "AppleGothic")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(5)


def add_docx_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("FITROOM 프로젝트 보고서   |   ")
    set_run_font(run, 8, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_docx_paragraph(doc: Document, text: str, *, bold_prefix: str | None = None, style=None) -> None:
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, bold=True, color=DARK)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)


def add_docx_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(item)
        set_run_font(run, 9.2)


def add_docx_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, DARK)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(value)
        set_run_font(run, 8.3, True, "FFFFFF")
        if widths:
            cell.width = Cm(widths[index])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cell, "F7F9FC")
            run = cell.paragraphs[0].add_run(value)
            set_run_font(run, 8.1)
            if widths:
                cell.width = Cm(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_docx_image(doc: Document, filename: str, caption: str) -> None:
    path = ASSET_DIR / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(6.45))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(7)
    run = cap.add_run(caption)
    set_run_font(run, 8, color=MUTED)


def build_docx() -> None:
    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_docx_styles(doc)
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    add_docx_page_number(section.footer.paragraphs[0])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(55)
    run = p.add_run("AI 고급 프로젝트")
    set_run_font(run, 10, True, BLUE)
    p = doc.add_paragraph(style="Title")
    p.add_run("FITROOM AI 광고 스튜디오\n프로젝트 보고서")
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(18)
    run = sub.add_run("의류 소상공인의 광고 제작과 가상 매장을 연결한 웹 서비스")
    set_run_font(run, 13, color=MUTED)
    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.autofit = False
    for i, (label, value) in enumerate([
        ("프로젝트 기간", "2026년 9월 16일 - 2026년 10월 1일"),
        ("최근 검증", "2026년 9월 29일"),
        ("서비스 주소", "fitroom-wardrobe.pigeon99999.chatgpt.site"),
        ("현재 상태", "핵심 기능·배포·실제 모델 평가 완료 / 사용자 최종 확인 대기"),
    ]):
        left, right = meta.rows[i].cells
        set_cell_shading(left, PALE)
        set_cell_margins(left, top=140, bottom=140)
        set_cell_margins(right, top=140, bottom=140)
        r1 = left.paragraphs[0].add_run(label)
        set_run_font(r1, 8.5, True, BLUE)
        r2 = right.paragraphs[0].add_run(value)
        set_run_font(r2, 8.5)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    doc.add_heading("요약", level=1)
    add_docx_paragraph(doc, "FITROOM은 광고 제작 인력과 예산이 부족한 의류 소상공인을 위한 웹 서비스다. 판매자가 상품 사실과 광고 조건을 입력하면 OpenAI gpt-5-mini가 상품 특징, 스타일링, 일상 장면의 세 관점으로 SNS 광고 문구 초안을 제안한다. 판매자는 결과를 직접 수정하고 TXT와 PNG 카드로 저장할 수 있다.")
    add_docx_paragraph(doc, "소상공인이 등록한 상품은 별도의 가상 매장으로도 이어진다. 이용자는 게시된 매장에 들어가 자신의 체형에 가까운 3D 아바타에 옷을 입혀보고, 판매자가 입력한 실측과 예상 여유를 확인한 뒤 실제 상점 링크로 이동한다.")
    add_docx_table(doc, ["핵심 지표", "현재 결과"], [
        ["자동 검사", "106개 통과"],
        ["품질 검사", "TypeScript, ESLint, 프로덕션 빌드 통과"],
        ["서비스 배포", "Sites 소유자 전용 비공개 배포 완료"],
        ["반응형 검증", "390px 화면에서 가로 넘침 없음"],
        ["모델 품질 평가", "정상 4건 통과 · 평균 19.1초 · 사실 보존 5.00/5"],
    ], [5.0, 11.5])
    add_docx_image(doc, "fitroom-landing.png", "그림 1  판매자와 이용자의 역할을 구분한 첫 화면")

    doc.add_heading("1 프로젝트 배경과 문제 정의", level=1)
    doc.add_heading("1.1 대상 사용자", level=2)
    add_docx_bullets(doc, [
        "온라인 또는 오프라인에서 의류를 판매하는 소상공인",
        "광고 문구를 전담하는 직원이나 외주 예산이 부족한 운영자",
        "전문 디자인 도구나 프롬프트 경험 없이 SNS 홍보 초안이 필요한 사용자",
    ])
    doc.add_heading("1.2 해결하려는 문제", level=2)
    add_docx_paragraph(doc, "신상품을 등록할 때마다 상품 소개와 SNS 게시물을 직접 작성해야 한다. 반복 작업에 시간이 들고 어떤 관점으로 소개할지 떠올리기 어렵다. FITROOM은 판매자가 알고 있는 사실 정보를 구조화하고, 서로 다른 세 관점의 초안을 비교·수정할 수 있게 한다.")
    doc.add_heading("1.3 성공 기준", level=2)
    add_docx_bullets(doc, [
        "상품과 광고 조건을 한 흐름에서 입력할 수 있다.",
        "실제 생성형 AI가 일정한 구조의 초안 3개를 반환한다.",
        "입력하지 않은 소재·성능·가격을 사실처럼 만들지 않도록 통제한다.",
        "판매자가 결과를 수정해 복사·TXT·PNG로 저장할 수 있다.",
        "입력 오류와 모델 연결 실패에서도 입력과 기존 결과를 보존한다.",
    ])

    doc.add_heading("2 범위와 주요 의사결정", level=1)
    add_docx_table(doc, ["구분", "내용"], [
        ["생성형 AI 핵심", "상품 정보 → 광고 조건 → 초안 3개 → 검토·편집 → 저장"],
        ["판매자 기능", "상품·가격·재고·실측 관리, 광고 제작, 가상 매장 제작, 계정 동기화"],
        ["이용자 기능", "가상 상점 거리, 3D 피팅, 실측 비교, 코디 보관함, 로컬 추천"],
        ["서버 기능", "ChatGPT 인증, D1 계정 저장, 게시 검증 공개 카탈로그"],
        ["제외 범위", "생성형 이미지, SNS 자동 게시, 외부 쇼핑몰 자동 동기화, 사진 기반 정확한 3D 의상 생성"],
    ], [3.5, 13.0])
    add_docx_paragraph(doc, "개인 프로젝트 기간 안에 완성 가능한 생성 파이프라인을 우선하고, 기존 3D 피팅룸을 소상공인 상품 노출 경험으로 확장했다. 생성형 AI 성과는 광고 문구 생성으로 한정하고 추천·핏 계산·3D 변형은 규칙 기반 기능으로 구분했다.")

    doc.add_heading("3 서비스 흐름과 시스템 구조", level=1)
    add_docx_table(doc, ["단계", "브라우저", "서버·모델"], [
        ["1 입력", "상품 사실, 광고 조건, 이미지 미리보기", "-"],
        ["2 검사", "필수값·숫자·길이 검사", "Zod 요청 검사, 32KB 제한"],
        ["3 생성", "진행 상태와 중복 요청 방지", "Responses API, gpt-5-mini, Structured Outputs"],
        ["4 후처리", "초안 선택과 직접 편집", "관점·길이·중복·해시태그 구조 검사"],
        ["5 저장", "복사, TXT, PNG, 평가 JSON", "원본 이미지와 체형 사진은 저장하지 않음"],
    ], [2.0, 7.2, 7.3])
    add_docx_paragraph(doc, "상품 이미지는 브라우저 미리보기와 PNG 광고 카드 합성에만 사용한다. AI 요청에는 상품 텍스트와 광고 조건만 포함하며 체형 사진, 체형 수치, 상품 이미지 원본은 전송하지 않는다.")

    doc.add_heading("4 기술 스택과 데이터 설계", level=1)
    add_docx_table(doc, ["기술", "역할"], [
        ["React 19, TypeScript, Vinext", "판매자·이용자 화면과 서버 라우트"],
        ["OpenAI Responses API", "gpt-5-mini 광고 문구 생성"],
        ["Zod", "브라우저 저장, 요청, 모델 응답 계약 검사"],
        ["Three.js, MakeHuman", "체형 변형 아바타와 참고 의상"],
        ["MediaPipe Pose Landmarker", "브라우저 안에서 정면·측면 포즈 지점 탐지"],
        ["Cloudflare D1, Drizzle", "계정별 판매자 작업 공간 저장"],
        ["Sites ChatGPT 인증", "서버 동기화 사용자의 계정 분리"],
        ["Node.js test runner, tsx", "계산·저장·오류 처리 자동 검사"],
    ], [5.0, 11.5])
    doc.add_heading("4.1 판매자 데이터 공개 경계", level=2)
    add_docx_paragraph(doc, "판매자 상품과 매장 설정은 브라우저에 먼저 저장한 뒤 사용자가 계정 동기화를 실행하면 D1에 저장한다. 공개 카탈로그 API는 게시 상태, 가격·재고·필수 실측, 매장 배치와 미리보기 확인을 다시 검사한다. 통과한 매장·상품만 반환하며 사용자 ID, 이메일, 초안 상품, 초안 상품의 진열 ID는 응답에서 제외한다.")

    doc.add_heading("5 생성형 AI 설계", level=1)
    doc.add_heading("5.1 모델 선정", level=2)
    add_docx_paragraph(doc, "교육 과정에서 허용된 모델 중 비용과 응답 속도를 고려해 gpt-5-mini를 시작 모델로 선택했다. Responses API의 Structured Outputs를 사용해 제목, 본문, CTA, 해시태그를 포함한 JSON 형식을 강제한다. 별도의 파인튜닝이나 LoRA 학습은 수행하지 않았다.")
    doc.add_heading("5.2 입력 전처리", level=2)
    add_docx_bullets(doc, [
        "문자열의 앞뒤 공백을 제거하고 코드 포인트 기준 최대 길이를 검사한다.",
        "상품 특징은 쉼표와 줄바꿈으로 나눠 1개 이상 5개 이하로 제한한다.",
        "비어 있는 선택값은 빈 문자열이나 0 대신 null로 정규화한다.",
        "할인 홍보를 선택하면 판매자가 확인한 할인율을 필수로 요구한다.",
        "정의하지 않은 요청·응답 필드를 거부한다.",
    ])
    doc.add_heading("5.3 프롬프트 안전 규칙", level=2)
    add_docx_bullets(doc, [
        "입력 JSON을 명령이 아닌 데이터로 취급한다.",
        "입력된 상품명·색상·특징·소재·가격·할인율만 사실로 단정한다.",
        "방수·보온·신축성·체형 보정·재고·배송 등 미입력 사실을 만들지 않는다.",
        "코디와 일상 장면은 제안형 문장으로 표현한다.",
        "상품 특징·스타일링·일상 장면의 세 관점과 중복되지 않는 제목을 만든다.",
    ])

    doc.add_heading("6 사용자 인터페이스", level=1)
    add_docx_paragraph(doc, "첫 화면에서 판매자와 이용자의 목적을 나누고 독립 URL로 이동한다. 판매자 센터는 대시보드, 상품 관리, 가상 매장, AI 광고 메뉴로 구성한다. 광고 화면은 PC에서 이미지·입력·결과를 세 열로 보여주고 좁은 화면에서는 두 열과 한 열로 재배치한다.")
    add_docx_image(doc, "fitroom-ad-studio.png", "그림 2  상품 정보와 광고 조건을 입력하는 AI 광고 스튜디오")
    add_docx_paragraph(doc, "생성 전, 생성 중, 입력 오류, 연결 오류, 형식 오류, 완료 후 편집, 이전 입력 기준 결과, 재생성 실패 상태를 서로 다른 문구와 버튼 상태로 구분했다. 결과가 없을 때 고정 예문을 실제 AI 결과처럼 표시하지 않는다.")
    add_docx_image(doc, "fitroom-virtual-street.png", "그림 3  서버 게시 매장을 포함한 이용자 가상 상점 거리")

    doc.add_heading("7 3D 가상 매장과 핏 계산", level=1)
    add_docx_bullets(doc, [
        "판매자가 게시한 매장 소개·테마·상품 배치를 이용자 상점 거리에 반영한다.",
        "이용자는 상품 카드 클릭·드래그 또는 버튼으로 모자·상의·하의를 입히고, 입은 상품이나 아바타의 의상을 눌러 벗는다.",
        "판매자 사진의 대표색과 상품명·특징에서 판별한 셔츠·카펜터 형태를 참고 의상에 반영한다.",
        "흰 배경에 평면 촬영한 상의는 사용자가 선택하면 앞면 무늬를 현재 브라우저 탭의 3D 상의에 적용한다.",
        "의류 단면을 둘레로 환산해 신체 둘레와의 차이를 보여준다.",
        "누락 실측, 신축성 허리, 조절형 모자처럼 직접 비교하기 어려운 항목은 판단을 보류한다.",
        "3D 의상은 실루엣 참고 표현이며 원단 물리나 실제 착용 정확도를 보장하지 않는다.",
    ])

    doc.add_heading("8 검증 결과", level=1)
    add_docx_table(doc, ["검사 범위", "검사 수", "결과"], [
        ["판매자 상품·매장·게시·사진·의상 판별", "30", "통과"],
        ["광고 계약·프롬프트·API·평가·등록 상품 변환", "27", "통과"],
        ["3D 피팅룸·카탈로그·추천·사진·마네킹 핏", "49", "통과"],
        ["전체", "106", "통과"],
    ], [11.0, 2.3, 3.2])
    add_docx_bullets(doc, [
        "TypeScript 타입 검사, ESLint, 프로덕션 빌드 통과",
        "판매자 저장 → D1 계정 저장 → 공개 카탈로그 → 이용자 상점 거리 → 진열 매장 입장 확인",
        "공개 API에서 사용자 식별자와 초안이 제외되는지 확인",
        "390px 화면에서 서버 매장 표시와 가로 넘침 없음 확인",
        "배포된 공개 카탈로그 API의 HTTP 200 응답 확인",
        "마네킹 핏 검사 10건에서 허용 노출률 이내인지 확인",
        "iCloud 밖의 새 폴더에서 설치·테스트·타입 검사·프로덕션 빌드 재현",
    ])
    add_docx_paragraph(doc, "자동 검사는 코드와 데이터 계약의 안정성을 보여주지만 광고 문구의 창의성, 사실성, 실제 판매 효과를 증명하지 않는다.")

    doc.add_heading("9 모델 및 서비스 평가 결과", level=1)
    add_docx_paragraph(doc, "2026년 9월 29일 활성 교육용 API 프로젝트에서 정상 입력 4건을 실제 생성하고 오류 입력 4건의 사전 거부를 확인했다. 정상 응답은 모두 구조 검사를 통과했고 평균 응답 시간은 19.1초, 범위는 14.8–22.5초였다. 미입력 성능 표현과 숫자 후보는 최종 자동 검사에서 0건이었다.")
    add_docx_table(doc, ["평가 항목", "판단 기준", "현재 상태"], [
        ["사실 보존", "입력하지 않은 상품 사실을 단정하지 않는가", "보조 검토 5.00/5"],
        ["톤·목적 반영", "고객층·말투·홍보 목적이 드러나는가", "보조 검토 3.75/5"],
        ["형식 일관성", "세 관점과 필수 필드를 유지하는가", "4/4건 통과"],
        ["응답 속도", "실제 작업 흐름에서 기다릴 수 있는가", "평균 19.1초"],
        ["활용 가능성", "적은 수정으로 게시 초안에 쓸 수 있는가", "보조 검토 4.00/5"],
    ], [3.2, 9.7, 3.6])
    add_docx_paragraph(doc, "첫 실행에서는 모델이 해시태그에 #과 구두점을 붙여 엄격 검증에서 거부됐다. 표시용 해시태그만 정규화한 뒤 기존 검사를 다시 적용해 해결했다. 보조 검토 점수는 사용자가 승인하거나 수정해야 사람 평가로 확정된다.")
    add_docx_paragraph(doc, "이 결과는 프롬프트 2026-09-16.1 기준이다. 12개 초안을 다시 읽어 판매자 요청 문장의 본문 인용, 신상 표현 누락, 초안 간 반복 문장·같은 마지막 안내, 항목 나열식 문장을 확인했고, 말투·목적·마지막 안내 값별 기준을 넣은 2026-09-30.1 판으로 프롬프트를 고쳤다. 자동 검사에도 같은 항목을 찾는 검사를 추가했다. 개정판을 실제 모델로 평가한 결과는 아직 없다.")

    doc.add_heading("10 개인정보와 안전 설계", level=1)
    add_docx_bullets(doc, [
        "OpenAI API 키는 서버 환경변수에서만 읽고 브라우저 코드와 Git에 포함하지 않는다.",
        "상품 사진은 AI 요청과 서버 저장에서 제외하고, 대표색과 앞면 무늬 처리는 브라우저에서만 수행한다.",
        "체형 사진은 브라우저 Worker에서 처리하고 원본을 저장하거나 전송하지 않는다.",
        "모델 요청의 저장 옵션을 비활성화한다.",
        "광고 요청 본문은 32KB로 제한하고 운영 로그에 상품 입력을 남기지 않는다.",
        "생성 결과는 게시 전 판매자 확인이 필요하다고 표시한다.",
        "공개 카탈로그에서 계정 식별자와 초안 데이터를 제외한다.",
    ])

    doc.add_heading("11 성과와 한계", level=1)
    add_docx_table(doc, ["성과", "근거"], [
        ["생성 파이프라인 구현", "입력·요청 검사·구조화 출력·편집·저장 흐름 완성"],
        ["서비스 역할 분리", "판매자 센터와 이용자 가상 매장의 독립 URL·화면"],
        ["계정과 공개 데이터 분리", "인증 D1 작업 공간과 게시 검증 공개 API"],
        ["설명 가능한 피팅", "단면·둘레 차이와 판단 보류 사유 표시"],
        ["재현 가능한 검증", "106개 자동 검사와 타입·린트·빌드 기록"],
    ], [5.0, 11.5])
    doc.add_heading("11.1 현재 한계", level=2)
    add_docx_bullets(doc, [
        "실제 API 생성은 평가했지만 사용자 확정 평가와 소상공인 사용성 테스트가 아직 없다.",
        "상품 이미지는 서버에 저장하지 않아 공개 매장에서 기본 3D 의상 이미지로 표시한다.",
        "단어 기반 의심 표현 검사는 문맥을 완전히 이해하지 못한다.",
        "PNG 광고 카드는 한 가지 레이아웃만 제공한다.",
        "3D 의상은 일부 색상·형태·상의 앞면 무늬만 반영하며 실제 원단·주름·착용 결과를 재현하지 않는다.",
        "상의 앞면 무늬는 현재 탭에서만 유지되고 다른 기기·이용자에게 전달되지 않는다.",
        "프로젝트 폴더가 iCloud 동기화 대상이면 큰 파일이 비워져 개발 서버와 빌드가 멈출 수 있다.",
    ])
    doc.add_heading("11.2 개선 순서", level=2)
    add_docx_bullets(doc, [
        "사용자가 생성 원문과 보조 점수를 확인해 사람 평가를 확정한다.",
        "처음 사용하는 사람의 생성 → 편집 → 저장 흐름을 관찰한다.",
        "상품 이미지 저장 정책과 광고 카드 템플릿 선택 기능을 설계한다.",
        "실제 휴대전화에서 터치·가독성·3D 성능을 확인한다.",
    ])

    doc.add_heading("12 재현 방법", level=1)
    add_docx_paragraph(doc, "Node.js 22.13 이상 환경에서 다음 명령을 실행한다.")
    code = doc.add_paragraph()
    for line in ["npm run install:ci", "npm test", "npm run typecheck", "npm run build", "npm run dev", "npm run evaluate:ads -- --dry-run"]:
        run = code.add_run(line + "\n")
        run.font.name = "Menlo"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Menlo")
        run.font.size = Pt(8.5)
    add_docx_paragraph(doc, "실제 모델 평가는 OPENAI_API_KEY가 서버 환경에 안전하게 설정된 경우에만 npm run evaluate:ads로 실행한다. 배포 사이트는 소유자 전용이며 ChatGPT 로그인이 필요하다.")

    doc.add_heading("13 AI 도구 지원 범위", level=1)
    add_docx_paragraph(doc, "서비스 기능의 생성 모델은 OpenAI gpt-5-mini다. MediaPipe Pose Landmarker는 사진의 포즈 지점을 찾는 사전 학습 모델이며, 예상 핏과 개인화 추천은 AI가 아닌 규칙·수식 기반 기능이다. 별도의 모델 학습이나 파인튜닝은 수행하지 않았다.")
    add_docx_paragraph(doc, "Codex의 GPT-6 기반 코딩 에이전트는 요구사항 정리, UI·서버 코드 작성, 테스트, 브라우저 검증, 문서 초안을 지원했다. 2026년 9월 29일에는 Claude Code(Claude Sonnet 5.5)가 결함 수정, 클릭 착용·탈착, 마네킹 핏 검사, 사진 대표색과 상의 앞면 무늬 처리, 문서 정리를 지원했다. 사용자는 서비스 방향, 대상 사용자, 가상 매장 경험, 일정과 교육 API 사용 가능 여부를 결정했다. 개인 업무일지의 배운 점과 소감은 사용자가 자신의 경험으로 보완한다.")

    doc.add_heading("14 제출 전 남은 작업", level=1)
    add_docx_bullets(doc, [
        "배포 환경에 OpenAI API 비밀값 연결과 실제 사이트 생성 확인",
        "처음 사용하는 사람의 생성·수정·저장 사용성 점검",
        "LMS에 비공개 GitHub 저장소(cjkj1234/fitroom, 2026-09-29 게시) 링크 제출",
        "업무일지의 개인 경험·배운 점·소감 보완",
    ])
    add_docx_paragraph(doc, "이 보고서는 현재 검증된 구현 결과와 남은 평가를 구분한다. 측정하지 않은 모델 품질과 실제 핏 정확도를 완료 성과로 주장하지 않는다.")

    doc.core_properties.title = "FITROOM AI 광고 스튜디오 프로젝트 보고서"
    doc.core_properties.subject = "의류 소상공인을 위한 생성형 AI 광고 제작 및 가상 매장 서비스"
    doc.core_properties.author = "FITROOM 프로젝트"
    doc.core_properties.keywords = "생성형 AI, OpenAI, 소상공인, 광고 문구, 3D 가상 매장"
    doc.save(DOCX_PATH)


def pdf_styles():
    pdfmetrics.registerFont(TTFont("AppleGothic", str(FONT_PATH)))
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("TitleKR", parent=base["Title"], fontName="AppleGothic", fontSize=28, leading=36, textColor=rl_color(DARK), alignment=TA_LEFT, spaceAfter=12, wordWrap="CJK"),
        "subtitle": ParagraphStyle("SubtitleKR", parent=base["Normal"], fontName="AppleGothic", fontSize=12, leading=19, textColor=rl_color(MUTED), wordWrap="CJK"),
        "h1": ParagraphStyle("H1KR", parent=base["Heading1"], fontName="AppleGothic", fontSize=18, leading=25, textColor=rl_color(DARK), spaceBefore=10, spaceAfter=8, keepWithNext=True, wordWrap="CJK"),
        "h2": ParagraphStyle("H2KR", parent=base["Heading2"], fontName="AppleGothic", fontSize=12, leading=18, textColor=rl_color(BLUE), spaceBefore=8, spaceAfter=5, keepWithNext=True, wordWrap="CJK"),
        "body": ParagraphStyle("BodyKR", parent=base["BodyText"], fontName="AppleGothic", fontSize=8.8, leading=14.2, textColor=rl_color(INK), spaceAfter=6, wordWrap="CJK"),
        "small": ParagraphStyle("SmallKR", parent=base["BodyText"], fontName="AppleGothic", fontSize=7.5, leading=11, textColor=rl_color(MUTED), wordWrap="CJK"),
        "bullet": ParagraphStyle("BulletKR", parent=base["BodyText"], fontName="AppleGothic", fontSize=8.5, leading=13.5, leftIndent=10, firstLineIndent=-7, textColor=rl_color(INK), spaceAfter=3, wordWrap="CJK"),
        "table": ParagraphStyle("TableKR", parent=base["BodyText"], fontName="AppleGothic", fontSize=7.3, leading=10.5, textColor=rl_color(INK), wordWrap="CJK"),
        "table_head": ParagraphStyle("TableHeadKR", parent=base["BodyText"], fontName="AppleGothic", fontSize=7.5, leading=10.5, textColor=colors.white, wordWrap="CJK"),
        "caption": ParagraphStyle("CaptionKR", parent=base["BodyText"], fontName="AppleGothic", fontSize=7.4, leading=11, textColor=rl_color(MUTED), alignment=TA_CENTER, spaceAfter=8, wordWrap="CJK"),
        "code": ParagraphStyle("CodeKR", parent=base["Code"], fontName="Courier", fontSize=7.7, leading=11, textColor=rl_color(DARK), leftIndent=8, spaceAfter=6),
    }


def p(text: str, style) -> Paragraph:
    return Paragraph(escape(text).replace("\n", "<br/>"), style)


def pdf_table(styles, headers: list[str], rows: list[list[str]], widths: list[float]) -> Table:
    data = [[p(item, styles["table_head"]) for item in headers]]
    data.extend([[p(item, styles["table"]) for item in row] for row in rows])
    table = Table(data, colWidths=[value * mm for value in widths], repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), rl_color(DARK)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.4, rl_color(LINE)),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, rl_color("F7F9FC")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    return table


def pdf_bullets(styles, items: list[str]) -> list[Paragraph]:
    return [p(f"- {item}", styles["bullet"]) for item in items]


def pdf_image(styles, filename: str, caption: str) -> list:
    path = ASSET_DIR / filename
    image = Image(str(path), width=170 * mm, height=106.25 * mm)
    return [image, Spacer(1, 3), p(caption, styles["caption"])]


def header_footer(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFont("AppleGothic", 7)
    canvas.setFillColor(rl_color(MUTED))
    canvas.drawString(20 * mm, 12 * mm, "FITROOM AI 광고 스튜디오 프로젝트 보고서")
    canvas.drawRightString(190 * mm, 12 * mm, str(doc.page))
    canvas.restoreState()


def first_page(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFillColor(rl_color(BLUE))
    canvas.rect(0, 267 * mm, 210 * mm, 30 * mm, fill=1, stroke=0)
    canvas.setFillColor(rl_color(DARK))
    canvas.rect(0, 0, 15 * mm, 267 * mm, fill=1, stroke=0)
    canvas.restoreState()


def build_pdf() -> None:
    PDF_PATH.parent.mkdir(parents=True, exist_ok=True)
    styles = pdf_styles()
    story = [Spacer(1, 35 * mm), p("AI 고급 프로젝트", styles["h2"]), p("FITROOM AI 광고 스튜디오\n프로젝트 보고서", styles["title"]), p("의류 소상공인의 광고 제작과 가상 매장을 연결한 웹 서비스", styles["subtitle"]), Spacer(1, 18 * mm)]
    story.append(pdf_table(styles, ["항목", "내용"], [
        ["프로젝트 기간", "2026년 9월 16일 - 2026년 10월 1일"],
        ["최근 검증", "2026년 9월 29일"],
        ["서비스 주소", "fitroom-wardrobe.pigeon99999.chatgpt.site"],
        ["현재 상태", "핵심 기능·배포·실제 모델 평가 완료 / 사용자 최종 확인 대기"],
    ], [35, 135]))
    story += [PageBreak(), p("요약", styles["h1"]), p("FITROOM은 광고 제작 인력과 예산이 부족한 의류 소상공인을 위한 웹 서비스다. 판매자가 상품 사실과 광고 조건을 입력하면 OpenAI gpt-5-mini가 상품 특징, 스타일링, 일상 장면의 세 관점으로 SNS 광고 문구 초안을 제안한다. 판매자는 결과를 직접 수정하고 TXT와 PNG 카드로 저장할 수 있다.", styles["body"]), p("소상공인이 등록한 상품은 별도의 가상 매장으로도 이어진다. 이용자는 게시된 매장에 들어가 자신의 체형에 가까운 3D 아바타에 옷을 입혀보고, 판매자가 입력한 실측과 예상 여유를 확인한 뒤 실제 상점 링크로 이동한다.", styles["body"])]
    story.append(pdf_table(styles, ["핵심 지표", "현재 결과"], [
        ["자동 검사", "106개 통과"], ["품질 검사", "TypeScript, ESLint, 프로덕션 빌드 통과"], ["서비스 배포", "Sites 소유자 전용 비공개 배포 완료"], ["반응형 검증", "390px 화면에서 가로 넘침 없음"], ["모델 품질 평가", "정상 4건 통과 · 평균 19.1초 · 사실 보존 5.00/5"],
    ], [43, 127]))
    story += [Spacer(1, 8)] + pdf_image(styles, "fitroom-landing.png", "그림 1  판매자와 이용자의 역할을 구분한 첫 화면")

    story += [PageBreak(), p("1 프로젝트 배경과 문제 정의", styles["h1"]), p("1.1 대상 사용자", styles["h2"])] + pdf_bullets(styles, ["온라인 또는 오프라인에서 의류를 판매하는 소상공인", "광고 문구를 전담하는 직원이나 외주 예산이 부족한 운영자", "전문 디자인 도구나 프롬프트 경험 없이 SNS 홍보 초안이 필요한 사용자"])
    story += [p("1.2 해결하려는 문제", styles["h2"]), p("신상품을 등록할 때마다 상품 소개와 SNS 게시물을 직접 작성해야 한다. 반복 작업에 시간이 들고 어떤 관점으로 소개할지 떠올리기 어렵다. FITROOM은 판매자가 알고 있는 사실 정보를 구조화하고 서로 다른 세 관점의 초안을 비교·수정할 수 있게 한다.", styles["body"]), p("1.3 성공 기준", styles["h2"])] + pdf_bullets(styles, ["상품과 광고 조건을 한 흐름에서 입력할 수 있다.", "실제 생성형 AI가 일정한 구조의 초안 3개를 반환한다.", "입력하지 않은 소재·성능·가격을 사실처럼 만들지 않도록 통제한다.", "판매자가 결과를 수정해 복사·TXT·PNG로 저장할 수 있다.", "입력 오류와 모델 연결 실패에서도 입력과 기존 결과를 보존한다."])
    story += [p("2 범위와 주요 의사결정", styles["h1"]), pdf_table(styles, ["구분", "내용"], [["생성형 AI 핵심", "상품 정보 → 광고 조건 → 초안 3개 → 검토·편집 → 저장"], ["판매자 기능", "상품·가격·재고·실측 관리, 광고 제작, 가상 매장 제작, 계정 동기화"], ["이용자 기능", "가상 상점 거리, 3D 피팅, 실측 비교, 코디 보관함, 로컬 추천"], ["서버 기능", "ChatGPT 인증, D1 계정 저장, 게시 검증 공개 카탈로그"], ["제외 범위", "생성형 이미지, SNS 자동 게시, 외부 쇼핑몰 자동 동기화, 사진 기반 정확한 3D 의상 생성"]], [34, 136]), p("개인 프로젝트 기간 안에 완성 가능한 생성 파이프라인을 우선하고 기존 3D 피팅룸을 소상공인 상품 노출 경험으로 확장했다. 생성형 AI 성과는 광고 문구 생성으로 한정하고 추천·핏 계산·3D 변형은 규칙 기반 기능으로 구분했다.", styles["body"])]

    story += [PageBreak(), p("3 서비스 흐름과 시스템 구조", styles["h1"]), pdf_table(styles, ["단계", "브라우저", "서버·모델"], [["1 입력", "상품 사실, 광고 조건, 이미지 미리보기", "-"], ["2 검사", "필수값·숫자·길이 검사", "Zod 요청 검사, 32KB 제한"], ["3 생성", "진행 상태와 중복 요청 방지", "Responses API, gpt-5-mini, Structured Outputs"], ["4 후처리", "초안 선택과 직접 편집", "관점·길이·중복·해시태그 구조 검사"], ["5 저장", "복사, TXT, PNG, 평가 JSON", "원본 이미지와 체형 사진은 저장하지 않음"]], [20, 72, 78]), p("상품 이미지는 브라우저 미리보기와 PNG 광고 카드 합성에만 사용한다. AI 요청에는 상품 텍스트와 광고 조건만 포함하며 체형 사진, 체형 수치, 상품 이미지 원본은 전송하지 않는다.", styles["body"]), p("4 기술 스택과 데이터 설계", styles["h1"]), pdf_table(styles, ["기술", "역할"], [["React 19, TypeScript, Vinext", "판매자·이용자 화면과 서버 라우트"], ["OpenAI Responses API", "gpt-5-mini 광고 문구 생성"], ["Zod", "브라우저 저장, 요청, 모델 응답 계약 검사"], ["Three.js, MakeHuman", "체형 변형 아바타와 참고 의상"], ["MediaPipe Pose Landmarker", "브라우저 안에서 정면·측면 포즈 지점 탐지"], ["Cloudflare D1, Drizzle", "계정별 판매자 작업 공간 저장"], ["Sites ChatGPT 인증", "서버 동기화 사용자의 계정 분리"], ["Node.js test runner, tsx", "계산·저장·오류 처리 자동 검사"]], [52, 118]), p("4.1 판매자 데이터 공개 경계", styles["h2"]), p("판매자 상품과 매장 설정은 브라우저에 먼저 저장한 뒤 사용자가 계정 동기화를 실행하면 D1에 저장한다. 공개 카탈로그 API는 게시 상태, 가격·재고·필수 실측, 매장 배치와 미리보기 확인을 다시 검사한다. 통과한 매장·상품만 반환하며 사용자 ID, 이메일, 초안 상품, 초안 상품의 진열 ID는 응답에서 제외한다.", styles["body"])]

    story += [PageBreak(), p("5 생성형 AI 설계", styles["h1"]), p("5.1 모델 선정", styles["h2"]), p("교육 과정에서 허용된 모델 중 비용과 응답 속도를 고려해 gpt-5-mini를 시작 모델로 선택했다. Responses API의 Structured Outputs를 사용해 제목, 본문, CTA, 해시태그를 포함한 JSON 형식을 강제한다. 별도의 파인튜닝이나 LoRA 학습은 수행하지 않았다.", styles["body"]), p("5.2 입력 전처리", styles["h2"])] + pdf_bullets(styles, ["문자열의 앞뒤 공백을 제거하고 코드 포인트 기준 최대 길이를 검사한다.", "상품 특징은 쉼표와 줄바꿈으로 나눠 1개 이상 5개 이하로 제한한다.", "비어 있는 선택값은 빈 문자열이나 0 대신 null로 정규화한다.", "할인 홍보를 선택하면 판매자가 확인한 할인율을 필수로 요구한다.", "정의하지 않은 요청·응답 필드를 거부한다."])
    story += [p("5.3 프롬프트 안전 규칙", styles["h2"])] + pdf_bullets(styles, ["입력 JSON을 명령이 아닌 데이터로 취급한다.", "입력된 상품명·색상·특징·소재·가격·할인율만 사실로 단정한다.", "방수·보온·신축성·체형 보정·재고·배송 등 미입력 사실을 만들지 않는다.", "코디와 일상 장면은 제안형 문장으로 표현한다.", "상품 특징·스타일링·일상 장면의 세 관점과 중복되지 않는 제목을 만든다."])
    story += [p("6 사용자 인터페이스", styles["h1"]), p("첫 화면에서 판매자와 이용자의 목적을 나누고 독립 URL로 이동한다. 판매자 센터는 대시보드, 상품 관리, 가상 매장, AI 광고 메뉴로 구성한다. 광고 화면은 PC에서 이미지·입력·결과를 세 열로 보여주고 좁은 화면에서는 두 열과 한 열로 재배치한다.", styles["body"])] + pdf_image(styles, "fitroom-ad-studio.png", "그림 2  상품 정보와 광고 조건을 입력하는 AI 광고 스튜디오")
    story += [p("생성 전, 생성 중, 입력 오류, 연결 오류, 형식 오류, 완료 후 편집, 이전 입력 기준 결과, 재생성 실패 상태를 서로 다른 문구와 버튼 상태로 구분했다. 결과가 없을 때 고정 예문을 실제 AI 결과처럼 표시하지 않는다.", styles["body"])]
    story += pdf_image(styles, "fitroom-virtual-street.png", "그림 3  서버 게시 매장을 포함한 이용자 가상 상점 거리")

    story += [PageBreak(), p("7 3D 가상 매장과 핏 계산", styles["h1"])] + pdf_bullets(styles, ["판매자가 게시한 매장 소개·테마·상품 배치를 이용자 상점 거리에 반영한다.", "이용자는 상품 카드 클릭·드래그 또는 버튼으로 옷을 입히고, 입은 상품이나 아바타의 의상을 눌러 벗는다.", "판매자 사진의 대표색과 상품명·특징에서 판별한 셔츠·카펜터 형태를 참고 의상에 반영한다.", "평면 촬영한 상의는 사용자가 선택하면 앞면 무늬를 현재 탭의 3D 상의에 적용한다.", "의류 단면을 둘레로 환산해 신체 둘레와의 차이를 보여준다.", "누락 실측, 신축성 허리, 조절형 모자는 판단을 보류한다.", "3D 의상은 실루엣 참고 표현이며 원단 물리나 실제 착용 정확도를 보장하지 않는다."])
    story += [p("8 검증 결과", styles["h1"]), pdf_table(styles, ["검사 범위", "검사 수", "결과"], [["판매자 상품·매장·게시·사진·의상 판별", "30", "통과"], ["광고 계약·프롬프트·API·평가·등록 상품 변환", "27", "통과"], ["3D 피팅룸·카탈로그·추천·사진·마네킹 핏", "49", "통과"], ["전체", "106", "통과"]], [105, 25, 40])] + pdf_bullets(styles, ["TypeScript 타입 검사, ESLint, 프로덕션 빌드 통과", "판매자 저장 → D1 계정 저장 → 공개 카탈로그 → 이용자 상점 거리 → 진열 매장 입장 확인", "공개 API에서 사용자 식별자와 초안이 제외되는지 확인", "390px 화면에서 서버 매장 표시와 가로 넘침 없음 확인", "마네킹 핏 검사 10건에서 허용 노출률 이내인지 확인", "iCloud 밖의 새 폴더에서 설치·테스트·타입 검사·프로덕션 빌드 재현"])
    story += [p("자동 검사는 코드와 데이터 계약의 안정성을 보여주지만 광고 문구의 창의성, 사실성, 실제 판매 효과를 증명하지 않는다.", styles["body"]), p("9 모델 및 서비스 평가 결과", styles["h1"]), p("2026년 9월 29일 정상 입력 4건을 실제 생성하고 오류 입력 4건의 사전 거부를 확인했다. 정상 응답은 모두 구조 검사를 통과했고 평균 응답 시간은 19.1초, 범위는 14.8–22.5초였다. 미입력 성능 표현과 숫자 후보는 최종 자동 검사에서 0건이었다.", styles["body"]), pdf_table(styles, ["평가 항목", "판단 기준", "현재 상태"], [["사실 보존", "입력하지 않은 상품 사실을 단정하지 않는가", "보조 검토 5.00/5"], ["톤·목적 반영", "고객층·말투·홍보 목적이 드러나는가", "보조 검토 3.75/5"], ["형식 일관성", "세 관점과 필수 필드를 유지하는가", "4/4건 통과"], ["응답 속도", "실제 작업 흐름에서 기다릴 수 있는가", "평균 19.1초"], ["활용 가능성", "적은 수정으로 게시 초안에 쓸 수 있는가", "보조 검토 4.00/5"]], [36, 98, 36]), p("첫 실행에서는 해시태그의 #과 구두점 때문에 엄격 검증에서 거부됐다. 표시용 해시태그만 정규화해 해결했으며, 보조 점수는 사용자가 승인하거나 수정해야 사람 평가로 확정된다.", styles["body"]), p("이 결과는 프롬프트 2026-09-16.1 기준이다. 12개 초안을 다시 읽어 판매자 요청 문장의 본문 인용, 신상 표현 누락, 초안 간 반복 문장·같은 마지막 안내, 항목 나열식 문장을 확인했고, 말투·목적·마지막 안내 값별 기준을 넣은 2026-09-30.1 판으로 프롬프트를 고쳤다. 자동 검사에도 같은 항목을 찾는 검사를 추가했다. 개정판을 실제 모델로 평가한 결과는 아직 없다.", styles["body"])]

    story += [p("10 개인정보와 안전 설계", styles["h1"])] + pdf_bullets(styles, ["OpenAI API 키는 서버 환경변수에서만 읽고 브라우저 코드와 Git에 포함하지 않는다.", "상품 사진은 AI 요청과 서버 저장에서 제외하고 대표색과 앞면 무늬 처리는 브라우저에서만 수행한다.", "체형 사진은 브라우저 Worker에서 처리하고 원본을 저장하거나 전송하지 않는다.", "모델 요청의 저장 옵션을 비활성화한다.", "광고 요청 본문은 32KB로 제한하고 운영 로그에 상품 입력을 남기지 않는다.", "생성 결과는 게시 전 판매자 확인이 필요하다고 표시한다.", "공개 카탈로그에서 계정 식별자와 초안 데이터를 제외한다."])
    story += [p("11 성과와 한계", styles["h1"]), pdf_table(styles, ["성과", "근거"], [["생성 파이프라인 구현", "입력·요청 검사·구조화 출력·편집·저장 흐름 완성"], ["서비스 역할 분리", "판매자 센터와 이용자 가상 매장의 독립 URL·화면"], ["계정과 공개 데이터 분리", "인증 D1 작업 공간과 게시 검증 공개 API"], ["설명 가능한 피팅", "단면·둘레 차이와 판단 보류 사유 표시"], ["재현 가능한 검증", "106개 자동 검사와 타입·린트·빌드 기록"]], [50, 120]), p("11.1 현재 한계", styles["h2"])] + pdf_bullets(styles, ["실제 API 생성은 평가했지만 사용자 확정 평가와 소상공인 사용성 테스트가 아직 없다.", "상품 이미지는 서버에 저장하지 않아 공개 매장에서 기본 3D 의상 이미지로 표시한다.", "단어 기반 의심 표현 검사는 문맥을 완전히 이해하지 못한다.", "PNG 광고 카드는 한 가지 레이아웃만 제공한다.", "3D 의상은 일부 색상·형태·상의 앞면 무늬만 반영하며 실제 원단·주름·착용 결과를 재현하지 않는다.", "상의 앞면 무늬는 현재 탭에서만 유지되고 다른 기기·이용자에게 전달되지 않는다.", "iCloud 동기화가 큰 파일을 비우면 개발 서버와 빌드가 멈출 수 있다."])
    story += [p("11.2 개선 순서", styles["h2"])] + pdf_bullets(styles, ["사용자가 생성 원문과 보조 점수를 확인해 사람 평가를 확정한다.", "처음 사용하는 사람의 생성 → 편집 → 저장 흐름을 관찰한다.", "상품 이미지 저장 정책과 광고 카드 템플릿 선택 기능을 설계한다.", "실제 휴대전화에서 터치·가독성·3D 성능을 확인한다."])

    story += [PageBreak(), p("12 재현 방법", styles["h1"]), p("Node.js 22.13 이상 환경에서 다음 명령을 실행한다.", styles["body"]), p("npm run install:ci\nnpm test\nnpm run typecheck\nnpm run build\nnpm run dev\nnpm run evaluate:ads -- --dry-run", styles["code"]), p("실제 모델 평가는 OPENAI_API_KEY가 서버 환경에 안전하게 설정된 경우에만 npm run evaluate:ads로 수행한다. 배포 사이트는 소유자 전용이며 ChatGPT 로그인이 필요하다.", styles["body"]), p("13 AI 도구 지원 범위", styles["h1"]), p("서비스 기능의 생성 모델은 OpenAI gpt-5-mini다. MediaPipe Pose Landmarker는 사진의 포즈 지점을 찾는 사전 학습 모델이며, 예상 핏과 개인화 추천은 AI가 아닌 규칙·수식 기반 기능이다. 별도의 모델 학습이나 파인튜닝은 수행하지 않았다.", styles["body"]), p("Codex의 GPT-6 기반 코딩 에이전트는 요구사항 정리, UI·서버 코드 작성, 테스트, 브라우저 검증, 문서 초안을 지원했다. 2026년 9월 29일에는 Claude Code(Claude Sonnet 5.5)가 결함 수정, 클릭 착용·탈착, 마네킹 핏 검사, 사진 대표색과 상의 앞면 무늬 처리, 문서 정리를 지원했다. 사용자는 서비스 방향과 일정, 교육 API 사용 가능 여부를 결정했다. 개인 업무일지의 배운 점과 소감은 사용자가 자신의 경험으로 보완한다.", styles["body"]), p("14 제출 전 남은 작업", styles["h1"])] + pdf_bullets(styles, ["배포 환경에 OpenAI API 비밀값 연결과 실제 사이트 생성 확인", "처음 사용하는 사람의 생성·수정·저장 사용성 점검", "LMS에 비공개 GitHub 저장소(cjkj1234/fitroom, 2026-09-29 게시) 링크 제출", "업무일지의 개인 경험·배운 점·소감 보완"])
    story += [p("이 보고서는 현재 검증된 구현 결과와 남은 평가를 구분한다. 측정하지 않은 모델 품질과 실제 핏 정확도를 완료 성과로 주장하지 않는다.", styles["body"])]

    doc = SimpleDocTemplate(str(PDF_PATH), pagesize=A4, rightMargin=20 * mm, leftMargin=20 * mm, topMargin=18 * mm, bottomMargin=19 * mm, title="FITROOM AI 광고 스튜디오 프로젝트 보고서", author="FITROOM 프로젝트")
    doc.build(story, onFirstPage=first_page, onLaterPages=header_footer)


if __name__ == "__main__":
    build_docx()
    build_pdf()
    print(DOCX_PATH)
    print(PDF_PATH)
